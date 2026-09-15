#!/usr/bin/env python3
"""Kalpi pitch D — dense intertwined deck."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

INK = RGBColor(0x1A, 0x1F, 0x1C)
PAPER = RGBColor(0xF4, 0xEF, 0xE4)
CIVIC = RGBColor(0x1F, 0x4F, 0x4A)
OUT = Path(__file__).with_name("kalpi-pitch-d.pptx")
ASSETS = Path(__file__).resolve().parents[2] / "design" / "assets"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
TOTAL = 22


def run(p, text, size, color=INK, font="Calibri", bold=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = font
    r.font.bold = bold


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER


def kicker(slide, text):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(0.22), Inches(12.2), Inches(0.32))
    run(box.text_frame.paragraphs[0], text.upper(), 11, CIVIC, "Calibri", True)


def title(slide, text, top=0.55, size=24, height=0.95):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(12.2), Inches(height))
    box.text_frame.word_wrap = True
    run(box.text_frame.paragraphs[0], text, size, INK, "Georgia")


def bullets(slide, lines, top=1.6, size=14, height=5.3):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(top), Inches(12.2), Inches(height))
    tf = box.text_frame
    tf.word_wrap = True
    for i, line in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        run(p, line, size, INK, "Calibri")


def footer(slide, n):
    box = slide.shapes.add_textbox(Inches(0.55), Inches(7.12), Inches(8), Inches(0.26))
    run(box.text_frame.paragraphs[0], "Kalpi  ·  pitch D", 11, CIVIC, "Calibri")
    box2 = slide.shapes.add_textbox(Inches(10.5), Inches(7.12), Inches(2.3), Inches(0.26))
    p = box2.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run(p, f"{n}  /  {TOTAL}", 11, CIVIC, "Calibri")


def new():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    return s


s = new()
kicker(s, "Pitch D  ·  intertwined  ·  strategy earns a product line  ·  cut later")
title(s, "Kalpi / קלפי", 0.7, 44, 1.0)
bullets(s, [
    "Free digital collect-only TCG of the 2026 lists. Pack = bait. Smart vote + show-up = goal.",
    "Game never picks a side. You may. Election 27 Oct. Lists 9 Sept. Threshold 3.25%. 120 / 61.",
    "No battles, lore, or IAP. Equal 8-card skeleton. Server pulls. Client displays.",
    "C = civic story. This deck = product a room can cut from. Notes are longer than the slides.",
], 1.85, 16)
footer(s, 1)

s = new()
kicker(s, "STRATEGY  ·  What we can prove  ·  why this cohort")
title(s, "46.7% of 18–24s   vs   74.9% of 65–74s   ·   ~600k first-timers", 0.52, 22, 0.85)
bullets(s, [
    "Knesset RIC after 2021. National 2022: 70.6%. A 28-point age gap — persistent, not a collapse.",
    "Lavi & Adler: youth vote more than young Europeans; still far below older Israelis.",
    "595–640k first-timers (18–22). ~8.7% of the roll. On the order of six seats. Points / IDI / CBS.",
    "Still the target: lowest turnout, largest new bloc, weakest list-knowledge. Older voters already show up.",
    "This group can move six seats without knowing slots 1–20. Lists 9 Sept. Election 27 Oct. Virality needs weeks.",
], 1.5, 15)
footer(s, 2)

s = new()
kicker(s, "STRATEGY  ·  Surveys are not turnout  ·  two jobs")
title(s, "80–95% intend. That is not a forecast.", 0.52, 26, 0.8)
bullets(s, [
    "IDI May 2026, n=576 first-timers: 80–95% intend.",
    "Same institute, 2013: ~88% intended, ~92% later said they voted, official 67.8%.",
    "18–22 already most likely to stay home (17.7% / ~15% after) vs 7–9% older. Overreport is known.",
    "Assume 2026 youth turnout stays low vs older adults — even if a few points above 46.7%.",
    "Two jobs: get them there, and make the slip informed. Feed is #1 source; most do not seek.",
    "No invented “X% cannot name three parties.” Say: they do not know the list.",
], 1.45, 15)
footer(s, 3)

s = new()
kicker(s, "STRATEGY + PRODUCT  ·  They do not know who  ·  so the product…")
title(s, "“Bennett is a bro.”  Lived. Not a fake knowledge %.", 0.52, 24, 0.75)
bullets(s, [
    "Me — a vibe is not a platform. Little brother — does not know. Hadar Muchtar — same room.",
    "NEW VOTERS → teach lists + the booth. Do not assume they know.",
    "PARENT DEFAULT → facts without a family fight. Not “rebel against home.”",
    "DON’T CARE (war, load) → 90-second rip. Dopamine first. Lecture never.",
    "TIKTOK / FAKE NEWS as civics → sourced backs + commons. Knowledge, not clip-by-clip.",
    "AFRAID TO SHOW AN OPINION → in-game faction is anonymous (tile, not a name).",
], 1.4, 15)
footer(s, 4)

s = new()
kicker(s, "STRATEGY  ·  Why not parents  ·  why literacy  ·  jam goals")
title(s, "Rebellion. They are invested. Solution: literacy.", 0.52, 24, 0.8)
bullets(s, [
    "A parent lecture is the wrong teacher. Parents already decided — that is a fight, not a syllabus.",
    "אוריינות פוליטית: teach a smart vote without a parent, a party ad, or a TikTok. They chose to flip.",
    "Fake news supports this: if they know the list, a fake has less room. We do not fact-check clip-by-clip.",
    "Jam: knowledge (backs/quiz/commons) · socials (share/tally/trade) · encourage voting (habit to 27 Oct).",
    "Taking a choice (faction, survey pack, then a real slip) · literacy = north star · young audience.",
], 1.45, 15)
footer(s, 5)

s = new()
kicker(s, "STRATEGY + PRODUCT  ·  Trait → this format, why now")
title(s, "If the product is a lecture, we already lost this age.", 0.52, 22, 0.75)
bullets(s, [
    "47s focus, median 40 (Gloria Mark) → rip under 90s. Not a PDF.",
    "Variable rewards → the rip is the habit engine, on purpose.",
    "Feed politics → lesson on the card back + commons + quiz.",
    "Afraid to speak → anonymous faction. Viral or invisible → share the 1/1.",
    "Pocket (~150M / ~$1.3B y1) → daily pack, ritual, binder, screenshot. Not battles.",
    "State Makers sold out a 98-card Israeli history set. Collectr: TCG sales +107% in 2025.",
    "Digital + free = reach, speed, no pay-to-win, no civic/money fight. Hit it now — lists 9 Sept.",
], 1.4, 14)
footer(s, 6)

s = new()
kicker(s, "STRATEGY + PRODUCT  ·  Local proof")
title(s, "State Makers / יוצרים מדינה — do not compress to “cards work.”", 0.52, 22, 0.85)
bullets(s, [
    "Piece of History. 98 cards. Packs of 8. Holos. First edition sold out.",
    "People paid to collect — and learned the country’s story. The format teaches here.",
    "We point that at this election: digital, free, equal lists, no till. Our pack is 6 — faster than their 8.",
    "They are cardboard / history-as-hobby. Next: why digital, and the global wave.",
    "pieceofhistory.com/statemakers",
], 1.55, 16)
footer(s, 7)

s = new()
kicker(s, "STRATEGY + PRODUCT  ·  Pocket’s five  ·  what we refuse")
title(s, "Daily pack. Ritual. Sets. Screenshot. Quick session.", 0.52, 22, 0.75)
bullets(s, [
    "1  Clock, not a pile — one free pack / 24h. (8h drops parked, more hostile, not Alpha.)",
    "2  Tear / pause — addictive on purpose. A fade-to-grid is not a collectible.",
    "3  Binder hole — collect sets, not a duel. Missing well = missing fact.",
    "4  Rarity you screenshot — that is the virality. Social is show, not chat.",
    "5  Under 90 seconds. Earlier digital Pokémon skipped collecting and stalled.",
    "Take the door. Refuse shop, hourglasses, Wonder-Pick paid rewind, battles.",
    "Why not a quiz app? Quiz apps are not a daily habit for this age. Pocket showed what is.",
], 1.4, 14)
footer(s, 8)

s = new()
kicker(s, "STRATEGY  ·  The iron is hot  ·  window is September")
title(s, "The format left Pokémon. Subject is not the limiter.", 0.52, 22, 0.75)
bullets(s, [
    "Pokémon still ~65% of tracked hobby. Pocket (30 Oct 2024) is the digital daily-pack proof.",
    "One Piece: nothing in 2022, hobby #2 (~17%) now. A new IP can become a collect giant fast.",
    "Riftbound (LoL, 31 Oct 2025): fastest 30-day move Cardlogs tracked. Cyberpunk TCG: ~$27M KS.",
    "Sports never left. Life TCG prints holos to animal population — we invert that to list slot.",
    "US politics cards exist (Decision 2024; Politic Packs). State Makers = best Israeli analog, not worldwide.",
    "Collectr: tracked sales +107% in 2025. Direction, not a random market-size $. Hit it now.",
], 1.4, 14)
footer(s, 9)

s = new()
kicker(s, "PRODUCT  ·  Therefore Kalpi  ·  bait vs goal")
title(s, "Pack is bait. Smart vote + show-up is the goal.", 0.52, 24, 0.75)
bullets(s, [
    "Free digital collect-only TCG of the lists. Teach them well enough that they know the list at the kalpi.",
    "Card back is ONE lesson. Same job: quiz, commons, party pack, trading a dupe, real-life party work.",
    "A rip with no learning beat fails. Extra pack = quiz. Not spend. Not in-game faction.",
    "No battles, lore, IAP, client-trusted pulls, live polls in our voice, “Kalpi recommends,” outing a camp, cash Alpha.",
    "Equal 8-card skeleton + 4 commons. PoC: Likud · Democrats · Ra'am (map spread) = 28 sourced cards.",
], 1.4, 14)
footer(s, 10)

s = new()
kicker(s, "PRODUCT  ·  Collectible musts  ·  knowledge on top")
title(s, "Essence of a collectible. Learning bound to progression.", 0.52, 22, 0.7)
bullets(s, [
    "Rip ritual → hold → flip the back. ~1.5s Alpha; 3s is the idea (too hostile to launch). Never skip the back as default.",
    "Timer 1/24h. No hourglass shop. Pity after 3 dry packs = new knowledge. Duplicates still show the back.",
    "Printed mix 3C / 2U / 1R-or-H. Hidden odds are a loot-box tell. Card lands in a hole = missing fact.",
    "Screenshot the 1/1. Date can caption a holo; slot-number is the chase. Extra pack = quiz on a pulled card.",
    "New lists after 9 Sept. Pulse later = CEC facts only — never horse-race polls.",
], 1.35, 14)
footer(s, 11)

s = new()
kicker(s, "PRODUCT  ·  Core loop  ·  pack rules  ·  Alpha")
title(s, "login → rip 6 → hold → flip back → binder → share / quiz / tally", 0.52, 18, 0.75)
bullets(s, [
    "Session: 5s to pack, 20–30s rip, 20–40s flip, 10–15s binder. Default under 90s. Deep browse optional.",
    "Printed mix. Pity. Prefer an incomplete set. Commons every other pack. Server creates; client displays.",
    "Alpha: guest, daily pack, rip, flip, binder %, share image. Server pulls from day one.",
    "Designed later: quiz, faction, volunteer bundle, survey pack, streaks, numbered holos, trading, boards.",
    "No paid rerolls. Pledge does not grant a pack. Energy does not exist.",
], 1.4, 14)
footer(s, 12)

s = new()
kicker(s, "PRODUCT  ·  Quiz on pulled cards")
title(s, "You only get asked what you already ripped. Prefer the last pack.", 0.52, 20, 0.8)
bullets(s, [
    "Popup: that same face + name. Swipe role / known-for (from the back). Swipe list. 3–4 options. 5:00.",
    "Both right → one extra pack. Wrong or timeout → show the back. One win per day. New popup later if they fail.",
    "Cannot farm a card you never pulled. Dismissible. Never blocks the daily pack. No “who is good.”",
    "Commons quiz 120 / 3.25% / closed list / CEC date the same way. Login = habit. Learning = bonus pack.",
], 1.5, 15)
footer(s, 13)

s = new()
kicker(s, "PRODUCT  ·  Three choice mechanics  ·  energy is dead")
title(s, "Do not mix these. There is no currency.", 0.52, 22, 0.7)
bullets(s, [
    "1  IN-GAME FACTION: pick a list. Each pack +1 on a public tally. Scoreboard, not a poll. No extra pack. Anonymous. Switch ok. Unfactioned still play. Tiles same size; order = CEC letter, not polls.",
    "2  REAL-LIFE PARTY WORK: verified HQ / campaign / kalpi. Best way to learn. Pays a PACK BUNDLE. Does not auto-join the faction.",
    "3  SURVEY → PARTY PACK: issue survey, published table, same questions. Pack of matched-list politicians only. Tool, not endorsement. Never hides the binder. Never auto-factions.",
], 1.4, 14, 5.4)
footer(s, 14)

s = new()
kicker(s, "PRODUCT  ·  List-slot numbered holos  ·  Life TCG invert")
title(s, "Print run means something real. We print to official list slot.", 0.52, 20, 0.75)
bullets(s, [
    "Slot 1 → 1/1. Slot 2 → 1/2 · 2/2. Slot 3 → 1/3 · 2/3 · 3/3. Slot n → 1/n … n/n.",
    "Likud 1, Democrats 1, Ra'am 1 are all 1/1. Teach order, not poll share. Never seats. Never “importance.”",
    "Life TCG prints holos to wild population. We steal the idea and invert it for neutrality.",
    "Server assigns the number. Symbol / platform / quote / commons are not slot-numbered.",
    "Separate: a holo can carry the date as a screenshot caption. Slot-number is the chase.",
], 1.4, 15)
footer(s, 15)

s = new()
kicker(s, "CONTENT  ·  What is on a card  ·  voice")
title(s, "Member. Symbol. Platform. Quote. Commons.", 0.52, 22, 0.7)
bullets(s, [
    "Per list: 4 members · 1 symbol · 2 platform · 1 quote. Same rarity math. Same art weight. No cool party.",
    "Commons: Knesset 120 · closed list · threshold 3.25% (since 2015) · CEC 9 September.",
    "Back: who / role / one vote-useful fact. Quote labeled + dated + sourced. Source on the card.",
    "Hebrew when we ship. Tired 19-year-old voice. No campaign “we.” No slogans as facts.",
    "Pulse later: CEC calendar only. Never live polls or “X is leading” in our voice. Civic objects, not press faces.",
], 1.35, 14)
footer(s, 16)

s = new()
kicker(s, "DESIGN  ·  Play it  ·  tokens")
x = 0.55
for name in ("hero-art-kalpi.png", "hero-art-threshold.png", "hero-art-memchetlammed.png"):
    path = ASSETS / name
    if path.exists():
        s.shapes.add_picture(str(path), Inches(x), Inches(1.05), Inches(2.0), Inches(2.8))
    x += 2.2
bullets(s, [
    "Paper, ink, foil. Party color = 8px pip only — never a wash. Shared chrome so no list is the hero.",
    "Date stamp on the holo. מחל = letters in the booth. Alpha: pack, flip, binder. Server pulls.",
    "Flip hero-cards.html  ·  phone wireframes  ·  visual system already locked.",
], 4.05, 15, 2.8)
footer(s, 17)

s = new()
kicker(s, "PRODUCT  ·  Why they stay  ·  metrics")
title(s, "Faction + cash are not enough.", 0.52, 22, 0.65)
bullets(s, [
    "ALPHA: daily clock, pity, binder %, share image.",
    "DESIGNED: quiz extra pack, faction tally (no reward), real-life bundle, survey party pack, trading (no paid gate),",
    "          slot 1/1 + date stamp, streaks 3/7, commons-complete badge, country board / leagues / set race.",
    "10k NIS first-set = pitch concept, legal later. Not Alpha.",
    "Metrics (deck, not Alpha): D1/D7 · % flipped to the back · quiz accuracy · extra packs from learning vs login · set %.",
    "Vanity (pulls, holos shared) only as leading indicators of the learning loop.",
], 1.25, 14)
footer(s, 18)

s = new()
kicker(s, "GROWTH  ·  How it leaves the room  ·  outbid invert")
title(s, "Share the rare. Seed where they are. Nobody pays for rank.", 0.52, 20, 0.75)
bullets(s, [
    "Screenshot the 1/1. TikTok / Instagram / WhatsApp. Influencers: their 18–24 is ours. Hebrew as soon as the rip works.",
    "Weavy / Figma Weave: stills, rip films, promo. Generate, then cut. Same deadline argument as AI art.",
    "outbid.lol (19 Aug 2026): pay-to-rank. Founder-reported ~1M visits, six-figure bids, top spot ~$13–17k.",
    "Proof people pay for a ladder. We invert — civic race, nobody pays. Treat dollars as founder-reported if picky.",
    "Lists 9 Sept. Virality needs weeks. Leave the room only if it can be live in September — free, Hebrew, seeded.",
], 1.4, 14)
footer(s, 19)

s = new()
kicker(s, "IMPLEMENTATION  ·  HTML first  ·  server owns luck")
title(s, "Open a URL. Server pulls. Client displays.", 0.52, 24, 0.7)
bullets(s, [
    "HTML first. No store. No framework tax. Tokens locked. Skills: frontend-design + kalpi-frontend.",
    "POST /packs/daily → clock / pity → RNG → persist inventory → six ids. Never POST { cards } as a pull.",
    "Quiz grade (owned ids), faction tally, volunteer bundle, party-pack match, streaks: server-side.",
    "Rate-limit: one daily pack, one quiz-pack per day. Guest ok; bind later.",
    "Weavy for motion. Civic objects, not scraped press faces. Artist pass if funded. Do not wait.",
], 1.35, 14)
footer(s, 20)

s = new()
kicker(s, "DESIGN  ·  The ends justify the means")
title(s, "Empty of lore. Full of cause.", 0.52, 26, 0.75)
bullets(s, [
    "Jams reward heart in the fiction. We shipped dopamine on purpose — rips, AI art, a faction board, no story bible.",
    "Civic goal is the message. Lists 9 Sept. Hand-painting every member of every list is how you miss October.",
    "Same psychology as gacha — and floors: no spend, short sessions, pity, extra packs only when they learn.",
    "Prior jams: guygir.itch.io — emptiness is a choice, not a missing soul.",
], 1.4, 15, 3.4)
qr = ASSETS / "qr-guygir-itchio.png"
if qr.exists():
    s.shapes.add_picture(str(qr), Inches(10.7), Inches(1.1), Inches(1.7), Inches(1.7))
footer(s, 21)

s = new()
kicker(s, "27 October 2026  ·  one line")
title(s, "Assume the gap. Give them a pack they will open.", 1.15, 28, 1.4)
bullets(s, [
    "Live in September. Free. Hebrew. Seeded where 18–24 already are.",
    "Civic story only: present C. This room was the product.",
    "Forbidden: IDI-as-forecast. Invented knowledge %. “Kalpi helps the Democrats.” Notes only.",
], 2.85, 16)
footer(s, 22)

prs.save(OUT)
print(f"Wrote {OUT}")
