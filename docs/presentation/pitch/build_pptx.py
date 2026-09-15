#!/usr/bin/env python3
"""Kalpi pitch: background, literacy, fit — one ask at the end."""

from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

INK = RGBColor(0x1A, 0x1F, 0x1C)
PAPER = RGBColor(0xF4, 0xEF, 0xE4)
CIVIC = RGBColor(0x1F, 0x4F, 0x4A)
OUT = Path(__file__).with_name("kalpi-pitch.pptx")
ASSETS = Path(__file__).resolve().parents[2] / "design" / "assets"

prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)
TOTAL = 23


def run(p, text, size, color=INK, font="Georgia", bold=False):
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.color.rgb = color
    r.font.name = font
    r.font.bold = bold


def bg(slide):
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = PAPER


def kicker(slide, text):
    box = slide.shapes.add_textbox(Inches(0.9), Inches(0.45), Inches(11.5), Inches(0.4))
    run(box.text_frame.paragraphs[0], text.upper(), 13, CIVIC, "Calibri", True)


def title(slide, text, top=1.2, size=40, height=2.5):
    box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.5), Inches(height))
    box.text_frame.word_wrap = True
    run(box.text_frame.paragraphs[0], text, size, INK, "Georgia")


def body(slide, text, top=3.9, size=22, height=2.5):
    box = slide.shapes.add_textbox(Inches(0.9), Inches(top), Inches(11.2), Inches(height))
    box.text_frame.word_wrap = True
    run(box.text_frame.paragraphs[0], text, size, INK, "Calibri")


def footer(slide, n):
    box = slide.shapes.add_textbox(Inches(0.9), Inches(7.05), Inches(6), Inches(0.3))
    run(box.text_frame.paragraphs[0], "Kalpi  ·  pitch", 12, CIVIC, "Calibri")
    box2 = slide.shapes.add_textbox(Inches(10.4), Inches(7.05), Inches(2), Inches(0.3))
    p = box2.text_frame.paragraphs[0]
    p.alignment = PP_ALIGN.RIGHT
    run(p, f"{n}  /  {TOTAL}", 12, CIVIC, "Calibri")


def new():
    s = prs.slides.add_slide(prs.slide_layouts[6])
    bg(s)
    return s


s = new()
kicker(s, "Israeli Elections 2026")
title(s, "Who does not reach the kalpi.", 1.3, 48)
body(s, "And why a pack of cards is the way to get them there — knowing who they are voting for.")
footer(s, 1)

s = new()
kicker(s, "1  ·  Voting %  ·  the background")
box = s.shapes.add_textbox(Inches(0.9), Inches(1.5), Inches(5.5), Inches(3.0))
tf = box.text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "46.7%", 72, CIVIC, "Georgia")
run(tf.add_paragraph(), "of 18–24s voted after 2021.", 20, INK, "Calibri")
box = s.shapes.add_textbox(Inches(7.0), Inches(1.5), Inches(5.5), Inches(3.0))
tf = box.text_frame
tf.word_wrap = True
run(tf.paragraphs[0], "74.9%", 72, CIVIC, "Georgia")
run(tf.add_paragraph(), "of 65–74s did.", 20, INK, "Calibri")
body(s, "Knesset Research and Information Center. National turnout in 2022: 70.6%.", 5.4, 16)
footer(s, 2)

s = new()
kicker(s, "Why that number matters")
title(s, "A 28-point age gap. Not a vibe.")
body(s, "Young Israelis vote more than young Europeans, and the youth rate has been rising. They still vote far less than older Israelis. That is the issue we are in the room for.")
footer(s, 3)

s = new()
kicker(s, "Why it matters this year")
box = s.shapes.add_textbox(Inches(0.9), Inches(1.4), Inches(11), Inches(1.5))
run(box.text_frame.paragraphs[0], "~600k", 80, CIVIC, "Georgia")
title(s, "first-timers. About six seats.", 3.1, 40, 1.2)
body(s, "595–640 thousand people aged 18–22. 8.7% of the roll. Election day 27 October 2026.", 4.6, 20)
footer(s, 4)

s = new()
kicker(s, "“But they say they will vote”")
title(s, "80–95% intend to. That is not a forecast.", 1.2, 36, 1.8)
body(s, "IDI, May 2026. In 2013 the same institute saw ~88% intend, ~92% later say they voted — official turnout was 67.8%. Ages 18–22 were already the most likely to stay home.", 3.4)
footer(s, 5)

s = new()
kicker(s, "Working assumption")
title(s, "Youth turnout will still be low.")
body(s, "Even if 2026 is a few points above 46.7%, it is still low versus older adults. And even the ones who show up often do not know the list. Two jobs: get them there, make the slip informed.")
footer(s, 6)

s = new()
kicker(s, "2  ·  They do not know who to vote for")
title(s, "“Bennett is a bro.”", 1.2, 44, 1.4)
body(s, "That was me. A vibe is not a platform. My little brother does not know who to vote for. Hadar Muchtar — same room. This is not a survey. This is the product.", 2.8)
footer(s, 7)

s = new()
kicker(s, "The young-voter problem")
title(s, "Default to parents. Or don’t care. Or the feed decides.", 1.2, 32, 2.2)
body(s, "War. Load. Enough already. TikTok is full of viral fake news. Some are afraid to show an opinion — at least on one side. A lot of new voters. Not a PDF audience.", 3.6)
footer(s, 8)

s = new()
kicker(s, "Why not “ask your parents”")
title(s, "Rebellion. And they are invested.")
body(s, "A parent lecture is the wrong teacher for this age. Parents are emotional and already decided. That is not a syllabus. That is a fight.")
footer(s, 9)

s = new()
kicker(s, "The solution")
title(s, "Political literacy. A smart vote.", 1.2, 40, 1.8)
body(s, "How do you teach them to vote properly — without a parent, a party ad, or a TikTok? Knowledge they chose to flip. Fake news is another reason this is the job: if they know the list, a fake has less room.", 3.3)
footer(s, 10)

s = new()
kicker(s, "Why they are still the target")
title(s, "Lowest turnout. Largest new bloc. Weakest list-knowledge.", 1.2, 34, 2.2)
body(s, "If anyone can move six seats without knowing who sits in slots 1–20, it is this group. Older voters already show up. This group does not — and when they do, the feed is their civics class.", 3.7)
footer(s, 11)

s = new()
kicker(s, "What characterizes them")
title(s, "47 seconds. Dopamine. Viral or invisible.")
body(s, "Gloria Mark: average screen focus is 47 seconds, median 40. Age of wars and TikTok. If it is not a daily loop, it is a lecture they will not finish.")
footer(s, 12)

s = new()
kicker(s, "3  ·  Local proof — cards already teach Israel")
title(s, "State Makers / יוצרים מדינה", 1.2, 40, 1.4)
body(s, "Piece of History. 98-card Israeli history set. Packs of 8. Holos. First edition sold out. People paid to collect — and learned the country’s story. We point that format at this election, digital and free.  pieceofhistory.com/statemakers", 2.8, 20)
footer(s, 13)

s = new()
kicker(s, "Why the digital maneuver works — Pokémon TCG Pocket")
title(s, "Five things we take. Not the shop.", 1.2, 32, 1.2)
body(s, "1  Daily timed pack — they come back.   2  The rip is a ritual — addictive.   3  Collecting sets is the progression.   4  Rarity is a screenshot — that is the virality.   5  Built for a quick session.", 2.6, 22, 3.2)
footer(s, 14)

s = new()
kicker(s, "STRATEGY  ·  The iron is hot")
title(s, "Not only Pokémon. The format is spreading.")
body(s, "Tracked TCG sales doubled in 2025. One Piece, Riftbound, Cyberpunk. Sports never left. US politics already has cards. State Makers is the best local proof. Hit it now.")
footer(s, 15)

s = new()
kicker(s, "Trait → game")
title(s, "90-second rip. Daily pack. Card back. Anonymous pledge. Share the pull.", 1.2, 28, 2.4)
body(s, "The rip is bait. Extra packs come from a 5-minute swipe. A survey can unlock a party pack. Show-up is the outcome.", 4.0, 20)
footer(s, 16)

s = new()
kicker(s, "They already pay for a ladder")
title(s, "outbid.lol — fame for whoever pays more.", 1.2, 34, 1.8)
body(s, "Companies spent five figures to sit on top of a public board. ~1M visits. Proof people will waste a lot of money to go viral. This time the spectacle is for a civic cause. Nobody pays.  outbid.lol  ·  outbid.lol/about", 3.3, 20)
footer(s, 17)

s = new()
kicker(s, "Therefore")
title(s, "Kalpi.", 1.4, 72, 1.6)
body(s, "A free digital collect-only TCG of the lists. No battles. No lore. No paywall. Equal sets. The game never picks a side — you may, anonymously. One free pack a day.")
footer(s, 18)

s = new()
kicker(s, "PRODUCT  ·  Bait vs goal")
title(s, "The pack is bait. The goal is a smart vote.")
body(s, "Teach Israeli politics well enough that they show up knowing the list. The card back is one lesson — not the only one. Quiz, commons, a party pack: same job. Progression is bound to learning.")
footer(s, 19)

s = new()
kicker(s, "Play it")
x = 0.9
for name in ("hero-art-kalpi.png", "hero-art-threshold.png", "hero-art-memchetlammed.png"):
    path = ASSETS / name
    if path.exists():
        s.shapes.add_picture(str(path), Inches(x), Inches(1.25), Inches(2.3), Inches(3.22))
    x += 2.55
body(s, "Alpha: pack, flip, binder. 28 sourced cards. Server pulls the pack — the player does not invent it.", 4.7, 20)
footer(s, 20)

s = new()
kicker(s, "Progression is learning")
title(s, "Image. Name. Two swipes. Five minutes.")
body(s, "What do they do? Which list? Both right → extra pack. Optional survey → a pack of that party’s politicians. Daily rip is habit. Literacy is how you get ahead.")
footer(s, 21)

s = new()
kicker(s, "The ends justify the means")
title(s, "Empty of lore. Full of cause.", 1.2, 40, 1.4)
body(s, "Jams usually reward heart in the fiction. This is viral on purpose. The civic goal is the message. Prior jams: guygir.itch.io", 2.8, 20, 1.6)
qr = ASSETS / "qr-guygir-itchio.png"
if qr.exists():
    s.shapes.add_picture(str(qr), Inches(10.4), Inches(1.3), Inches(2.0), Inches(2.0))
footer(s, 22)

s = new()
kicker(s, "27 October 2026")
title(s, "Assume the gap. Give them a pack they will open.")
body(s, "If this should leave the room, it needs to be live in September — free, Hebrew, seeded where 18–24 already are.")
footer(s, 23)

prs.save(OUT)
print(f"Wrote {OUT}")
